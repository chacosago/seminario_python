# -necesito saber qué alumnos de la lista de aweb no estan en la lista de siu
# listas ubicadas en el mismo directorio que el programa
# -2do quienes están en siu pero no en au

# La idea seria sacar una lista de diccionarios, cada dicc es un alumno del cual 
# guardo: legajo, apellido, nombre, mail.

import xlrd
import sys
import os
import openpyxl

#caracteres_tildes = "áàéèíìóòúù"
#caracteres_tildes_l = list(caracteres_tildes)

# descomentar en vscode:
#os.chdir("/home/hcastorp/Drive/Facultad/Informatica/Python/TP/otras-pruebas/proyecto-carla/")
print("Tenga en cuenta que los archivos a utilizar deben estar en la misma carpeta que el programa y deben incluir su extensión.\n")
if len(sys.argv) > 1:
    archivo_aweb = sys.argv[1]
    archivo_siu = sys.argv[2]
else:
    archivo_aweb = input("Ingresar nombre del arhivo descargado de AulasWeb: ")
    archivo_siu = input("Ingresar nombre del arhivo descargado de Siu-Guarani: ")
#excel_siu = xlrd.open_workbook('Inscripciones a cursadas.xls')
excel_siu = xlrd.open_workbook(archivo_siu)
hoja_siu = excel_siu.sheet_by_index(0)

alumnos_siu = []

for i in range(11,hoja_siu.nrows):
    fila = hoja_siu.row(i)
    apellido_nombre = hoja_siu.cell_value(i,1).split(",")
    alumnos_siu.append({"Legajo": hoja_siu.cell_value(i,0), "Apellido": apellido_nombre[0].lower(),
        "Nombre": apellido_nombre[1].lower(), "Mail": hoja_siu.cell_value(i,4)})

# Ahora con los de AulasWeb

#excel_auweb = openpyxl.load_workbook('TV2_COM2_2021  ELECCION DE GRUPOS.xlsx')
excel_auweb = openpyxl.load_workbook(archivo_aweb)
hoja_auweb = excel_auweb.active

alumnos_auweb = []

for fila in hoja_auweb.iter_rows(values_only=True):
    alumnos_auweb.append({"Apellido": fila[0].lower(), "Nombre": fila[1].lower(), "Mail": fila[3], "Comision" : fila[4]})

alumnos_no_encontrados = []

opcion = int(input("Ingrese la opcion deseada:\n1 - Para busacar alumnos que estan en Aulas Web pero no en Siu-Guarani.\n2 - Para busacar alumnos que estan en Siu-Guarani pero no en Aulas Web: "))

if opcion == 1:
    # Resuelvo problema: qué alumnos de la lista de aweb no estan en la lista de siu.
    for alumno_w in alumnos_auweb:
        # me fijo si esta en siu:
        for alumno_s in alumnos_siu:
            esta = alumno_w["Mail"] == alumno_s["Mail"] or ( alumno_w["Nombre"] in alumno_s["Nombre"] and alumno_w["Apellido"] in alumno_s["Apellido"])
            if esta:
                break
        if not esta:
            alumnos_no_encontrados.append(alumno_w)
    # Descarto la 1er fila con encabezados
    alumnos_no_encontrados.pop(0)
elif opcion == 2:
    # Resuelvo problema: qué alumnos de la lista de Siu no estan en la lista de Aweb.
    for alumno_s in alumnos_siu:
        # me fijo si esta en AuWeb:
        for alumno_w in alumnos_auweb:
            esta = alumno_s["Mail"] == alumno_w["Mail"] or ( alumno_w["Nombre"] in alumno_s["Nombre"] and alumno_w["Apellido"] in alumno_s["Apellido"])
            if esta:
                break
        if not esta:
            alumnos_no_encontrados.append(alumno_s)
else:
    print("Opcion invalida")



lista_ordenada = list(sorted(alumnos_no_encontrados,key=lambda t : t["Apellido"]))
c = 0

print("\nTener en cuenta que puede haber falsos positivos por tildes o simbolos\n")
for i in lista_ordenada:
    c+=1
    print(f'{c:3d} {i["Nombre"].capitalize():<30}{i.get("Apellido").capitalize():<30}{i.get("Mail"):<30}')

# para que en Windows no cierre la ventana al terminar la ejecucion
input("")